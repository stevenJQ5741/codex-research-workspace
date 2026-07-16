from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


EMU_PER_INCH = 914_400


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", " ", text).strip()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def resolve(base: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else (base / path).resolve()


def docx_cell_text(cell: Any) -> list[str]:
    chunks = [paragraph.text for paragraph in cell.paragraphs]
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                chunks.extend(docx_cell_text(nested_cell))
    return chunks


def inspect_docx(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(docx_cell_text(cell))
    structure = {
        "inline_shapes": len(document.inline_shapes),
        "sections": len(document.sections),
        "top_level_tables": len(document.tables),
    }
    return normalize("\n".join(chunks)), structure, []


def walk_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def shape_text(shape: Any) -> list[str]:
    chunks: list[str] = []
    if getattr(shape, "has_text_frame", False):
        chunks.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return chunks


def inspect_pptx(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    presentation = Presentation(path)
    slide_texts: list[str] = []
    pictures = 0
    charts = 0
    tables = 0
    shape_count = 0
    for slide in presentation.slides:
        chunks: list[str] = []
        for shape in walk_shapes(slide.shapes):
            shape_count += 1
            chunks.extend(shape_text(shape))
            pictures += int(shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
            charts += int(bool(getattr(shape, "has_chart", False)))
            tables += int(bool(getattr(shape, "has_table", False)))
        slide_texts.append(normalize("\n".join(chunks)))
    structure = {
        "slides": len(presentation.slides),
        "slide_width_inches": presentation.slide_width / EMU_PER_INCH,
        "slide_height_inches": presentation.slide_height / EMU_PER_INCH,
        "shapes": shape_count,
        "pictures": pictures,
        "charts": charts,
        "tables": tables,
    }
    return normalize("\n".join(slide_texts)), structure, slide_texts


def token_coverage(source: str, rendered: str) -> float:
    source_tokens = set(re.findall(r"\w{4,}", source.casefold(), flags=re.UNICODE))
    if not source_tokens:
        return 1.0
    rendered_tokens = set(re.findall(r"\w{4,}", rendered.casefold(), flags=re.UNICODE))
    return len(source_tokens & rendered_tokens) / len(source_tokens)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Validate a DOCX/PDF or PPTX/PDF research-brief release pair."
    )
    parser.add_argument("--spec", type=Path, required=True, help="UTF-8 JSON release spec")
    args = parser.parse_args()

    spec_path = args.spec.resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    base = spec_path.parent
    source_value = spec.get("source") or spec.get("docx") or spec.get("pptx")
    if not source_value:
        raise SystemExit("Spec must define source, docx, or pptx.")
    source_path = resolve(base, source_value)
    pdf_path = resolve(base, spec["pdf"])
    report_path = resolve(base, spec.get("report", "_work/qa_release.json"))

    missing = [str(path) for path in (source_path, pdf_path) if not path.exists()]
    if missing:
        raise SystemExit("Missing release artifact(s): " + ", ".join(missing))

    source_kind = source_path.suffix.lower().lstrip(".")
    if source_kind in {"docx", "docm", "dotx", "dotm"}:
        source_text, structure, source_units = inspect_docx(source_path)
        source_family = "docx"
    elif source_kind in {"pptx", "pptm", "ppsx", "ppsm", "potx", "potm"}:
        source_text, structure, source_units = inspect_pptx(source_path)
        source_family = "pptx"
    else:
        raise SystemExit(f"Unsupported Office source: {source_path}")

    reader = PdfReader(str(pdf_path))
    pdf_page_text = [normalize(page.extract_text() or "") for page in reader.pages]
    pdf_text = normalize("\n".join(pdf_page_text))

    required = [normalize(item) for item in spec.get("required_strings", [])]
    forbidden = [normalize(item) for item in spec.get("forbidden_strings", [])]
    required_case_sensitive = bool(spec.get("required_case_sensitive", True))

    def has(haystack: str, needle: str, case_sensitive: bool) -> bool:
        if case_sensitive:
            return needle in haystack
        return needle.casefold() in haystack.casefold()

    required_source = {
        item: has(source_text, item, required_case_sensitive) for item in required
    }
    required_pdf = {item: has(pdf_text, item, required_case_sensitive) for item in required}
    forbidden_source = {item: has(source_text, item, False) for item in forbidden}
    forbidden_pdf = {item: has(pdf_text, item, False) for item in forbidden}

    page_sizes = [
        [float(page.mediabox.width), float(page.mediabox.height)] for page in reader.pages
    ]
    expected_size = spec.get("expected_page_size_points")
    tolerance = float(spec.get("page_size_tolerance_points", 1.0))
    page_size_ok = True
    if expected_size:
        expected_width, expected_height = map(float, expected_size)
        page_size_ok = all(
            abs(width - expected_width) <= tolerance
            and abs(height - expected_height) <= tolerance
            for width, height in page_sizes
        )

    source_hash_results: dict[str, dict[str, Any]] = {}
    for source_name, expected_hash in spec.get("source_hashes", {}).items():
        protected_path = resolve(base, source_name)
        actual_hash = sha256(protected_path) if protected_path.exists() else None
        source_hash_results[source_name] = {
            "exists": protected_path.exists(),
            "expected": str(expected_hash).upper(),
            "actual": actual_hash,
            "matches": actual_hash == str(expected_hash).upper(),
        }

    minimum_chars = int(spec.get("minimum_text_characters_per_page", 1))
    checks: dict[str, bool] = {
        "expected_page_count": len(reader.pages) == int(spec["expected_pages"]),
        "expected_page_size": page_size_ok,
        "no_blank_or_nearly_blank_pages": all(
            len(text) >= minimum_chars for text in pdf_page_text
        ),
        "all_required_in_office_source": all(required_source.values()),
        "all_required_in_pdf": all(required_pdf.values()),
        "all_forbidden_absent_from_office_source": not any(forbidden_source.values()),
        "all_forbidden_absent_from_pdf": not any(forbidden_pdf.values()),
        "source_hashes_match": all(
            item["matches"] for item in source_hash_results.values()
        ),
    }

    count_mappings = {
        "expected_inline_shapes": "inline_shapes",
        "expected_sections": "sections",
        "expected_top_level_tables": "top_level_tables",
        "expected_slides": "slides",
        "expected_shapes": "shapes",
        "expected_pictures": "pictures",
        "expected_charts": "charts",
        "expected_tables": "tables",
    }
    for spec_key, structure_key in count_mappings.items():
        if spec_key in spec:
            checks[spec_key] = structure.get(structure_key) == int(spec[spec_key])

    if source_family == "pptx":
        checks["slide_count_matches_pdf_pages"] = structure["slides"] == len(reader.pages)
        if "expected_slide_size_inches" in spec:
            expected_width, expected_height = map(float, spec["expected_slide_size_inches"])
            slide_tolerance = float(spec.get("slide_size_tolerance_inches", 0.01))
            checks["expected_slide_size"] = (
                abs(structure["slide_width_inches"] - expected_width) <= slide_tolerance
                and abs(structure["slide_height_inches"] - expected_height) <= slide_tolerance
            )

    coverage: list[float] = []
    if source_family == "pptx" and source_units:
        coverage = [
            token_coverage(slide_text, pdf_page_text[index])
            for index, slide_text in enumerate(source_units)
            if index < len(pdf_page_text)
        ]
        if "minimum_slide_text_token_coverage" in spec:
            threshold = float(spec["minimum_slide_text_token_coverage"])
            checks["minimum_slide_text_token_coverage"] = (
                len(coverage) == len(source_units) and all(value >= threshold for value in coverage)
            )

    report = {
        "spec": str(spec_path),
        "source": str(source_path),
        "source_kind": source_family,
        "source_sha256": sha256(source_path),
        "pdf": str(pdf_path),
        "pdf_sha256": sha256(pdf_path),
        "page_count": len(reader.pages),
        "page_sizes_points": page_sizes,
        "page_text_characters": [len(text) for text in pdf_page_text],
        "source_structure": structure,
        "slide_text_token_coverage": coverage,
        "required_in_office_source": required_source,
        "required_in_pdf": required_pdf,
        "forbidden_in_office_source": forbidden_source,
        "forbidden_in_pdf": forbidden_pdf,
        "protected_source_hashes": source_hash_results,
        "checks": checks,
        "all_checks_pass": all(checks.values()),
        "visual_inspection_required": True,
    }

    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=True, indent=2))
    if not report["all_checks_pass"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
